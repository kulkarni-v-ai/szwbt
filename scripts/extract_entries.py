import openpyxl
import json

wb = openpyxl.load_workbook('Final list of entries.xlsx')
sheet = wb['Sheet3']
entries = []

for r in range(1, 103):
    num = sheet.cell(r, 1).value
    name = sheet.cell(r, 2).value
    state = sheet.cell(r, 3).value
    if name is not None:
        name = str(name).strip().replace('\ufffd', "'")
    if state is not None:
        state = str(state).strip()
    entries.append({
        'num': int(num) if num is not None else r,
        'name': name or '',
        'state': state or ''
    })

print(f"Total extracted: {len(entries)}")
with open('entries_extracted.json', 'w', encoding='utf-8') as f:
    json.dump(entries, f, indent=2, ensure_ascii=False)
print("Saved to entries_extracted.json successfully")
